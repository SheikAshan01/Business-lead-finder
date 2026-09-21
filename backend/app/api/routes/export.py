import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import or_, select
from sqlalchemy.orm import Session
from app.db import get_db
from app.models.business import Business, Category, LeadStatus, WebsiteStatus

router = APIRouter(prefix="/export", tags=["export"])

EXCEL_COLUMNS = [
    "Business Name",
    "Client Name",
    "Category",
    "Phone",
    "Email",
    "Address",
    "Area",
    "Taluk",
    "District",
    "Pincode",
    "Website",
    "Website Status",
    "Facebook",
    "Instagram",
    "WhatsApp",
    "Lead Score",
    "Lead Status",
    "Source",
    "Source URL",
]


def _build_export_query(
    website_status: WebsiteStatus | None = None,
    district: str | None = None,
    category_id: int | None = None,
    saved: bool | None = None,
    min_score: int | None = None,
    lead_status: LeadStatus | None = None,
    selected_ids: str | None = None,
):
    query = select(Business)
    if selected_ids:
        ids = [int(i.strip()) for i in selected_ids.split(",") if i.strip().isdigit()]
        if ids:
            return query.where(Business.id.in_(ids))

    if website_status:
        query = query.where(Business.website_status == website_status)
    if district and district.strip():
        query = query.where(Business.district.ilike(f"%{district.strip()}%"))
    if category_id:
        query = query.where(Business.category_id == category_id)
    if saved is not None:
        query = query.where(Business.is_saved == saved)
    if min_score is not None:
        query = query.where(Business.lead_score >= min_score)
    if lead_status:
        query = query.where(Business.lead_status == lead_status)

    return query


def _extract_row_data(b: Business) -> list:
    cat_name = b.category_rel.name if b.category_rel else ""
    web_status = b.website_status.value if hasattr(b.website_status, "value") else str(b.website_status)
    lead_status = b.lead_status.value if hasattr(b.lead_status, "value") else str(b.lead_status)

    return [
        b.business_name or "",
        b.client_name or "",
        cat_name,
        b.phone or "",
        b.email or "",
        b.address or "",
        b.area or "",
        b.taluk or "",
        b.district or "",
        b.pincode or "",
        b.website_url or "",
        web_status,
        b.facebook_url or "",
        b.instagram_url or "",
        b.whatsapp_url or "",
        b.lead_score,
        lead_status,
        b.source or "",
        b.source_url or "",
    ]


@router.get("/csv")
def export_csv(
    website_status: WebsiteStatus | None = None,
    district: str | None = None,
    category_id: int | None = None,
    saved: bool | None = None,
    min_score: int | None = None,
    lead_status: LeadStatus | None = None,
    selected_ids: str | None = None,
    db: Session = Depends(get_db),
):
    query = _build_export_query(
        website_status=website_status,
        district=district,
        category_id=category_id,
        saved=saved,
        min_score=min_score,
        lead_status=lead_status,
        selected_ids=selected_ids,
    )
    records = db.scalars(query.order_by(Business.lead_score.desc(), Business.id.desc())).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(EXCEL_COLUMNS)
    for b in records:
        writer.writerow(_extract_row_data(b))

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=sra-business-leads.csv"},
    )


@router.get("/excel")
def export_excel(
    website_status: WebsiteStatus | None = None,
    district: str | None = None,
    category_id: int | None = None,
    saved: bool | None = None,
    min_score: int | None = None,
    lead_status: LeadStatus | None = None,
    selected_ids: str | None = None,
    db: Session = Depends(get_db),
):
    query = _build_export_query(
        website_status=website_status,
        district=district,
        category_id=category_id,
        saved=saved,
        min_score=min_score,
        lead_status=lead_status,
        selected_ids=selected_ids,
    )
    records = db.scalars(query.order_by(Business.lead_score.desc(), Business.id.desc())).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "SRA Leads"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="101A2B", end_color="101A2B", fill_type="solid")
    ws.append(EXCEL_COLUMNS)

    for col_num in range(1, len(EXCEL_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for b in records:
        ws.append(_extract_row_data(b))

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=sra-business-leads.xlsx"},
    )
